import os
import re
import time
import uuid
import shutil
import logging
from typing import Dict, List, Any, Optional, Tuple

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
try:
    from webdriver_manager.chrome import ChromeDriverManager
except Exception:
    ChromeDriverManager = None

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Ensure Selenium cache path has write access in container/sandbox
os.environ.setdefault("SELENIUM_CACHE_DIR", "/tmp/.selenium")
os.makedirs("/tmp/.selenium", exist_ok=True)

logger = logging.getLogger("student_scraper")
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")

DOSSIERS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "generated_dossiers")
os.makedirs(DOSSIERS_DIR, exist_ok=True)

def cleanup_browser_processes():
    """Aggressively terminates any lingering chromium/chromedriver processes to reclaim RAM."""
    try:
        import subprocess
        subprocess.run(["pkill", "-9", "-f", "chromium|chromedriver"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except Exception:
        pass

def get_chrome_options(headless: bool = True) -> webdriver.ChromeOptions:
    """Configures Chrome options with extreme memory optimization (<512MB RAM) and anti-bot flags."""
    options = webdriver.ChromeOptions()
    if headless:
        options.add_argument("--headless=new")

    # Critical Low-RAM Flags for 512MB Free/Starter containers:
    options.add_argument("--no-sandbox")
    options.add_argument("--disable-dev-shm-usage")
    options.add_argument("--disable-gpu")
    options.add_argument("--disable-software-rasterizer")
    options.add_argument("--renderer-process-limit=1")
    options.add_argument("--no-zygote")
    options.add_argument("--window-size=960,600")

    # Cap V8 JS heap to 128MB max to trigger early GC
    options.add_argument("--js-flags=--max-old-space-size=128")

    # Disable image loading & decoding (saves >50% browser RAM)
    options.add_argument("--blink-settings=imagesEnabled=false")

    # Disable disk & media caches to prevent RAM buildup
    options.add_argument("--disk-cache-size=1")
    options.add_argument("--media-cache-size=1")
    options.add_argument("--aggressive-cache-discard")

    # Disable unnecessary background features & services
    options.add_argument("--disable-extensions")
    options.add_argument("--disable-default-apps")
    options.add_argument("--disable-sync")
    options.add_argument("--disable-translate")
    options.add_argument("--disable-background-networking")
    options.add_argument("--disable-background-timer-throttling")
    options.add_argument("--disable-backgrounding-occluded-windows")
    options.add_argument("--disable-breakpad")
    options.add_argument("--disable-component-update")
    options.add_argument("--disable-domain-reliability")
    options.add_argument("--disable-features=AudioServiceOutOfProcess,IsolateOrigins,site-per-process,TranslateUI,BlinkGenPropertyTrees")

    options.add_argument("--disable-infobars")
    options.add_argument("--disable-notifications")
    options.add_argument("--disable-blink-features=AutomationControlled")
    options.add_experimental_option("excludeSwitches", ["enable-automation"])
    options.add_experimental_option("useAutomationExtension", False)
    options.add_argument("user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36")

    # Detect system Chromium binary on Linux/Docker/Render
    for bin_path in [
        "/usr/bin/chromium",
        "/usr/bin/chromium-browser",
        "/usr/bin/google-chrome",
        "/usr/bin/google-chrome-stable",
    ]:
        if os.path.exists(bin_path):
            options.binary_location = bin_path
            logger.info(f"Using system browser binary at: {bin_path}")
            break

    return options

def create_driver(headless: bool = True) -> webdriver.Chrome:
    """Creates a Selenium Chrome driver with automatic fallback from system drivers to ChromeDriverManager."""
    options = get_chrome_options(headless=headless)

    # 1. First check system chromedriver (Docker / Render / Linux apt package)
    for driver_path in ["/usr/bin/chromedriver", "/usr/local/bin/chromedriver"]:
        if os.path.exists(driver_path):
            try:
                logger.info(f"Attempting system chromedriver at {driver_path}")
                service = Service(executable_path=driver_path)
                driver = webdriver.Chrome(service=service, options=options)
                driver.execute_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
                return driver
            except Exception as e:
                logger.warning(f"Failed using system chromedriver: {e}")

    # 2. Try ChromeDriverManager (for Mac / Windows / local dev)
    if ChromeDriverManager is not None:
        try:
            logger.info("Installing/verifying chromedriver via ChromeDriverManager...")
            service = Service(ChromeDriverManager().install())
            driver = webdriver.Chrome(service=service, options=options)
            driver.execute_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
            return driver
        except Exception as e:
            logger.warning(f"ChromeDriverManager failed: {e}. Falling back to native driver.")

    # 3. Fallback to native Selenium 4.6+ driver
    logger.info("Initializing native Selenium Chrome driver...")
    driver = webdriver.Chrome(options=options)
    driver.execute_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
    return driver


def generate_styled_excel(all_sections_data: List[Dict[str, str]], enrollment_no: str) -> str:
    """
    Generates a beautifully formatted Excel dossier matching the user's exact specification.
    Returns the absolute path to the generated file.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Details"
    ws.views.sheetView[0].showGridLines = True

    color_navy_deep = "0F172A"
    color_section_bar = "1E3A8A"
    color_sub_header = "2563EB"
    color_zebra_light = "F8FAFC"
    color_border_subtle = "CBD5E1"

    font_main_title = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    font_section_title = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_col_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_data_bold = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
    font_data_regular = Font(name="Segoe UI", size=10, color="0F172A")

    fill_navy = PatternFill(start_color=color_navy_deep, end_color=color_navy_deep, fill_type="solid")
    fill_sec = PatternFill(start_color=color_section_bar, end_color=color_section_bar, fill_type="solid")
    fill_col = PatternFill(start_color=color_sub_header, end_color=color_sub_header, fill_type="solid")
    fill_zebra = PatternFill(start_color=color_zebra_light, end_color=color_zebra_light, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border_side = Side(border_style="thin", color=color_border_subtle)
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 1. Main Title Banner (Row 2)
    ws.merge_cells("B2:D2")
    title_cell = ws["B2"]
    title_cell.value = f"STUDENT PROFILE & ACADEMIC DOSSIER — {enrollment_no}"
    title_cell.font = font_main_title
    title_cell.fill = fill_navy
    title_cell.alignment = align_center
    ws.row_dimensions[2].height = 40

    # 2. Table Column Headers (Row 4)
    col_headers = ["S.No.", "Field / Information", "Details / Value"]
    ws.row_dimensions[4].height = 25
    for col_idx, h_text in enumerate(col_headers, start=2):
        c = ws.cell(row=4, column=col_idx, value=h_text)
        c.font = font_col_header
        c.fill = fill_col
        c.alignment = align_center
        c.border = cell_border

    current_row = 5
    item_counter = 1
    current_sec = None

    for item in all_sections_data:
        sec = item["Section"]
        f_name = item["Field Name"]
        f_val = item["Value"]

        if sec != current_sec:
            current_sec = sec
            ws.row_dimensions[current_row].height = 26
            ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=4)
            sec_cell = ws.cell(row=current_row, column=2, value=f"  ◆  {sec.upper()}")
            sec_cell.font = font_section_title
            sec_cell.fill = fill_sec
            sec_cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1

        ws.row_dimensions[current_row].height = 20
        row_fill = fill_zebra if item_counter % 2 == 0 else fill_white

        c1 = ws.cell(row=current_row, column=2, value=item_counter)
        c1.font = font_data_bold
        c1.alignment = align_center
        c1.fill = row_fill
        c1.border = cell_border

        c2 = ws.cell(row=current_row, column=3, value=f_name)
        c2.font = font_data_bold
        c2.alignment = align_left
        c2.fill = row_fill
        c2.border = cell_border

        c3 = ws.cell(row=current_row, column=4, value=f_val)
        c3.font = font_data_regular
        c3.alignment = align_left
        c3.fill = row_fill
        c3.border = cell_border

        item_counter += 1
        current_row += 1

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 38
    ws.column_dimensions["D"].width = 58

    clean_enrollment = re.sub(r'[^a-zA-Z0-9_-]', '', enrollment_no) or "student"
    unique_suffix = uuid.uuid4().hex[:6]
    excel_filename = f"Student_Dossier_{clean_enrollment}_{unique_suffix}.xlsx"
    full_path = os.path.join(DOSSIERS_DIR, excel_filename)
    wb.save(full_path)
    try:
        wb.close()
    except Exception:
        pass
    del ws
    del wb
    import gc
    gc.collect()
    logger.info(f"Excel dossier saved at: {full_path}")
    return full_path


def extract_pumis_portal(driver, log) -> Tuple[List[Dict[str, str]], str]:
    """Extracts student details from modern student.pumis.in portal."""
    all_data = []
    greeting_name = ""

    # Extract student greeting from Dashboard
    try:
        body_text = driver.find_element(By.TAG_NAME, "body").text
        m_name = re.search(r'Welcome,\s*([A-Za-z\s]+)!', body_text)
        if m_name:
            greeting_name = m_name.group(1).strip()
            all_data.append({
                "Section": "Personal Information",
                "Field Name": "Student Name",
                "Value": greeting_name
            })
            log(f"Extracted student name from Dashboard greeting: {greeting_name}")
    except Exception:
        pass

    # Navigate to Profile
    log("Navigating to https://student.pumis.in/Profile...")
    driver.get("https://student.pumis.in/Profile")
    time.sleep(3.5)

    try:
        WebDriverWait(driver, 8).until(EC.presence_of_element_located((By.CLASS_NAME, "MuiTabPanel-root")))
    except Exception:
        time.sleep(2)

    p_body = driver.find_element(By.TAG_NAME, "body").text
    p_lines = [l.strip() for l in p_body.splitlines() if l.strip()]

    current_section = "Personal Information"
    skip_exact = {"dashboard", "profile", "search", "view", "edit", "logout", "home"}

    i = 0
    while i < len(p_lines) - 1:
        line = p_lines[i]
        line_lower = line.lower()
        if line_lower in skip_exact or line.startswith("http") or len(line) > 50:
            i += 1
            continue
        if line in ["Personal Details", "Basic Details"]:
            current_section = "Personal Information"
            i += 1
            continue
        elif line in ["Admission Details", "Academic Details"]:
            current_section = "Admission Information"
            i += 1
            continue
        elif line in ["Contact Information", "Address Details"]:
            current_section = "Contact Information"
            i += 1
            continue
        elif line in ["Parent/Guardian Details", "Parents Information"]:
            current_section = "Parents Information"
            i += 1
            continue
        elif line in ["Identification Information", "Identity Documents"]:
            current_section = "Identity Information"
            i += 1
            continue
        elif line in ["Campus Service Details", "Campus Services"]:
            current_section = "Campus Services"
            i += 1
            continue

        val = p_lines[i + 1]
        if not line.startswith("·") and val != line:
            if not any(sw in val.lower() for sw in ["personal details", "admission details", "contact information"]):
                all_data.append({
                    "Section": current_section,
                    "Field Name": line,
                    "Value": val
                })
                i += 2
                continue
        i += 1

    log(f"Extracted {len(all_data)} fields from modern Student Portal (pumis.in)")
    return all_data, greeting_name


def extract_classic_ums_portal(driver, wait, log) -> List[Dict[str, str]]:
    """Extracts student details from classic ums.paruluniversity.ac.in portal."""
    all_sections_data = []

    log("Opening Student Profile page on classic UMS...")
    profile_links = [
        "//a[contains(@class, 'user-profile') or contains(@class, 'dropdown-toggle')]",
        "//a[translate(normalize-space(text()), 'PROFILE', 'profile')='profile']",
        "//a[contains(@href, 'Profile') or contains(@href, 'StudentDetail')]",
        "//span[contains(text(), 'Profile')]/.."
    ]
    for plink in profile_links:
        try:
            p_elem = driver.find_element(By.XPATH, plink)
            if p_elem.is_displayed():
                driver.execute_script("arguments[0].click();", p_elem)
                time.sleep(1.2)
        except Exception:
            continue

    try:
        my_prof = driver.find_element(By.XPATH, "//a[contains(text(), 'Profile') or contains(@href, 'Profile')]")
        driver.execute_script("arguments[0].click();", my_prof)
        time.sleep(3.5)
    except Exception:
        pass

    sections_workflow = [
        {"name": "Personal Information", "keywords": ["personal", "basic"]},
        {"name": "Admission Information", "keywords": ["admission", "program"]},
        {"name": "Contact Information", "keywords": ["contact", "address"]},
        {"name": "Parents Information", "keywords": ["parent", "father", "mother"]},
        {"name": "Education Qualification", "keywords": ["education", "qualification 2", "qualification"]}
    ]

    def extract_from_active_tab():
        records = []
        seen = set()

        pane = None
        try:
            pane = driver.find_element(By.XPATH, "//div[contains(@class, 'tab-pane') and (contains(@class, 'active') or contains(@class, 'in'))]")
        except Exception:
            try:
                pane = driver.find_element(By.XPATH, "//div[contains(@class, 'tab-content')]")
            except Exception:
                pane = driver.find_element(By.TAG_NAME, "body")

        # 1. Inputs / Selects / Textareas (Form values)
        form_elements = pane.find_elements(By.XPATH, ".//input | .//select | .//textarea")
        for fe in form_elements:
            try:
                val = fe.get_attribute("value") or ""
                val = val.strip()
                if not val or val in ["-", "--", "null", "None"]:
                    continue

                label_text = ""
                elem_id = fe.get_attribute("id") or ""
                if elem_id:
                    try:
                        lbl = pane.find_element(By.XPATH, f".//label[@for='{elem_id}']")
                        label_text = lbl.text.strip()
                    except Exception:
                        pass

                if not label_text:
                    label_text = fe.get_attribute("name") or elem_id

                if label_text and val:
                    records.append((label_text, val))
            except Exception:
                continue

        # 2. Table row key-value pairs
        tables = pane.find_elements(By.XPATH, ".//table")
        for tbl in tables:
            try:
                rows = tbl.find_elements(By.XPATH, ".//tr")
                for r in rows:
                    cols = r.find_elements(By.XPATH, ".//td | .//th")
                    if len(cols) == 2:
                        k = cols[0].text.strip().replace(":", "").strip()
                        v = cols[1].text.strip()
                        if k and v:
                            records.append((k, v))
                    elif len(cols) >= 4:
                        k1 = cols[0].text.strip().replace(":", "").strip()
                        v1 = cols[1].text.strip()
                        k2 = cols[2].text.strip().replace(":", "").strip()
                        v2 = cols[3].text.strip()
                        if k1 and v1:
                            records.append((k1, v1))
                        if k2 and v2:
                            records.append((k2, v2))
            except Exception:
                continue

        # 3. Label-to-Span / Col pairs (Bootstrap grid)
        label_spans = pane.find_elements(By.XPATH, ".//label")
        for lbl in label_spans:
            try:
                k = lbl.text.strip().replace(":", "").strip()
                if not k or len(k) > 40:
                    continue
                try:
                    val_el = lbl.find_element(By.XPATH, "./following-sibling::*[1]")
                    v = val_el.text.strip() or val_el.get_attribute("value") or ""
                    if v:
                        records.append((k, v))
                except Exception:
                    pass
            except Exception:
                continue

        # 4. Text line-by-line fallback
        raw_text = pane.text
        lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
        i = 0
        while i < len(lines):
            line = lines[i]
            if ":" in line and not line.startswith("http"):
                parts = line.split(":", 1)
                k = parts[0].strip()
                v = parts[1].strip()
                if not v and (i + 1) < len(lines):
                    next_l = lines[i + 1]
                    if ":" not in next_l:
                        v = next_l
                        i += 1
                if k and v:
                    records.append((k, v))
            i += 1

        # Cleaning & Sanitization
        cleaned = []
        for k, v in records:
            k_clean = k.strip().replace("\n", " ").replace(":", "").strip()
            v_clean = v.strip().replace("\n", " ")

            if not v_clean or v_clean in ["-", "--", "N/A", "NA", "null", "None", ":"]:
                continue

            if any(fw in k_clean.lower() for fw in ["fee", "fees", "collected", "installment", "amount"]):
                continue

            combo = f"{k_clean.lower()}::{v_clean.lower()}"
            if combo not in seen:
                seen.add(combo)
                cleaned.append((k_clean, v_clean))

        return cleaned

    for sec in sections_workflow:
        sec_name = sec["name"]
        keywords = sec["keywords"]
        log(f"Extracting section: {sec_name}...")

        all_tabs = driver.find_elements(
            By.XPATH,
            "//a[@data-toggle='tab' or @role='tab'] | //ul[contains(@class, 'nav-tabs')]//a | //div[contains(@class, 'accordion')]//a | //div[contains(@class, 'card-header')]//a | //li[contains(@class, 'tab')]//a"
        )

        for t in all_tabs:
            try:
                t_txt = t.text.strip().lower()
                if any(kw in t_txt for kw in keywords):
                    if "fee" in t_txt:
                        continue
                    driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", t)
                    driver.execute_script("arguments[0].click();", t)
                    time.sleep(2.0)
                    break
            except Exception:
                continue

        items = extract_from_active_tab()
        log(f"Extracted {len(items)} fields for {sec_name}")

        for field_name, field_val in items:
            all_sections_data.append({
                "Section": sec_name,
                "Field Name": field_name,
                "Value": field_val
            })

    return all_sections_data


def scrape_student_data(enrollment_no: str, dob: str, headless: bool = True, log_cb=None) -> Dict[str, Any]:
    """
    Executes the complete headless scraping pipeline for University UMS.
    Returns:
      {
        "success": bool,
        "enrollment_no": str,
        "dob": str,
        "excel_path": str,
        "filename": str,
        "total_fields": int,
        "sections": { "Section Name": [{"field": k, "value": v}] },
        "summary": { "name": ..., "program": ..., ... },
        "error": Optional[str]
      }
    """
    def log(msg: str):
        logger.info(msg)
        if log_cb:
            try:
                log_cb(msg)
            except Exception:
                pass

    portal_url = "https://ums.paruluniversity.ac.in/StudentPanel/"
    all_sections_data: List[Dict[str, str]] = []
    driver = None

    try:
        # Reclaim any stray memory before launching browser
        cleanup_browser_processes()
        import gc
        gc.collect()

        log(f"Initializing invisible browser for {enrollment_no}...")
        driver = create_driver(headless=headless)
        wait = WebDriverWait(driver, 18)

        log("Navigating to UMS portal...")
        driver.get(portal_url)
        time.sleep(2)

        # ---------------------------------------------------------------------
        # STEP 1: STUDENT SECTION SWITCH
        # ---------------------------------------------------------------------
        log("Switching to Student section...")
        student_selectors = [
            "//input[@id='rblRole_1' or (@type='radio' and contains(@value, 'Student'))]",
            "//label[contains(text(), 'Student')]",
            "//a[translate(normalize-space(text()), 'STUDENT', 'student')='student']",
            "//button[translate(normalize-space(text()), 'STUDENT', 'student')='student']"
        ]
        for selector in student_selectors:
            try:
                elements = driver.find_elements(By.XPATH, selector)
                for el in elements:
                    if el.is_displayed():
                        driver.execute_script("arguments[0].click();", el)
                        time.sleep(1.2)
                        break
            except Exception:
                continue

        # ---------------------------------------------------------------------
        # STEP 2: ENROLLMENT & DOB FILLING (Sanitized to DDMMYYYY digits)
        # ---------------------------------------------------------------------
        # Parul University UMS strictly expects 8 digits DDMMYYYY for student birthdate password
        clean_dob = re.sub(r'[^0-9]', '', dob)
        pwd_to_send = clean_dob if len(clean_dob) == 8 else dob
        log(f"Entering enrollment number: {enrollment_no} and password/DOB: {pwd_to_send} (length {len(pwd_to_send)})...")

        enroll_input = wait.until(EC.visibility_of_element_located((
            By.XPATH, "//input[contains(@id, 'txtEnrollment') or contains(@id, 'txtUsername') or contains(@placeholder, 'Enrollment') or @type='text'][1]"
        )))
        enroll_input.clear()
        enroll_input.send_keys(enrollment_no)
        time.sleep(0.8)

        # Fill password / DOB directly or fallback via TAB
        pwd_filled = False
        try:
            pwd_input = driver.find_element(By.XPATH, "//input[contains(@id, 'txtPassword') or contains(@id, 'txtDob') or @type='password']")
            if pwd_input.is_displayed():
                pwd_input.clear()
                pwd_input.send_keys(pwd_to_send)
                pwd_filled = True
        except Exception:
            pass

        if not pwd_filled:
            enroll_input.send_keys(Keys.TAB)
            time.sleep(0.5)
            driver.switch_to.active_element.send_keys(pwd_to_send)
        time.sleep(0.8)

        # ---------------------------------------------------------------------
        # STEP 3: LOGIN BUTTON CLICK
        # ---------------------------------------------------------------------
        log("Submitting login credentials...")
        login_btn_selectors = [
            "//input[contains(@id, 'btnSubmit') or contains(@id, 'btnLogin')]",
            "//button[contains(@id, 'btnSubmit') or contains(@id, 'btnLogin')]",
            "//input[@type='submit' and contains(@value, 'Login')]"
        ]
        login_clicked = False
        for btn_sel in login_btn_selectors:
            try:
                btns = driver.find_elements(By.XPATH, btn_sel)
                for b in btns:
                    if b.is_displayed():
                        driver.execute_script("arguments[0].click();", b)
                        login_clicked = True
                        break
                if login_clicked:
                    break
            except Exception:
                continue

        if not login_clicked:
            enroll_input.send_keys(Keys.ENTER)

        # ---------------------------------------------------------------------
        # STEP 4: POPUP & REDIRECT HANDLING (Redirect to New Student Panel / Alerts)
        # ---------------------------------------------------------------------
        log("Checking for alerts and notification popups...")
        time.sleep(2)
        for _ in range(8):
            try:
                alert = driver.switch_to.alert
                log(f"Dismissed alert: {alert.text}")
                alert.dismiss()
                break
            except Exception:
                pass

            # Detect and click "Yes" on "Do you want to redirect to new student panel?"
            try:
                yes_btns = driver.find_elements(By.XPATH, "//button[contains(text(), 'Yes')] | //a[contains(text(), 'Yes')]")
                for yb in yes_btns:
                    if yb.is_displayed():
                        log("Detected 'Redirect to new student panel' prompt. Clicking 'Yes'...")
                        driver.execute_script("arguments[0].click();", yb)
                        time.sleep(3)
                        break
            except Exception:
                pass
            time.sleep(0.5)

        log("Waiting for Student Dashboard to load...")
        time.sleep(4)

        # Check if login was rejected
        curr_url = driver.current_url.lower()
        page_source = driver.page_source.lower()
        if "invalid" in page_source and ("password" in page_source or "username" in page_source or "enrollment" in page_source):
            raise ValueError("Authentication Failed: UMS portal reported invalid Enrollment Number or DOB.")

        # ---------------------------------------------------------------------
        # STEP 5: NAVIGATE TO PROFILE PAGE & HARVEST DATA
        # ---------------------------------------------------------------------
        greeting_name = ""
        is_pumis = "pumis.in" in driver.current_url.lower()

        if is_pumis:
            pumis_data, greeting_name = extract_pumis_portal(driver, log)
            all_sections_data.extend(pumis_data)
        else:
            classic_data = extract_classic_ums_portal(driver, wait, log)
            all_sections_data.extend(classic_data)


        if not all_sections_data:
            raise RuntimeError("No records found in student profile tabs. Please check credentials or portal layout.")

        # ---------------------------------------------------------------------
        # STEP 7: BUILD EXCEL FILE & RESPONSE STRUCTURE
        # ---------------------------------------------------------------------
        log("Building styled Excel dossier...")
        excel_path = generate_styled_excel(all_sections_data, enrollment_no)
        filename = os.path.basename(excel_path)

        # Build clean summary
        summary = {
            "enrollment_no": enrollment_no,
            "dob": dob,
            "student_name": greeting_name if greeting_name else "N/A",
            "course": "N/A",
            "institute": "N/A",
            "contact_mobile": "N/A",
            "email": "N/A",
            "domicile": "N/A"
        }

        sections_grouped: Dict[str, List[Dict[str, str]]] = {}
        for row in all_sections_data:
            s_name = row["Section"]
            f_key = row["Field Name"]
            f_val = row["Value"]

            if s_name not in sections_grouped:
                sections_grouped[s_name] = []
            sections_grouped[s_name].append({"field": f_key, "value": f_val})

            f_lower = f_key.lower()
            if "name" in f_lower and summary["student_name"] == "N/A" and "father" not in f_lower and "mother" not in f_lower:
                summary["student_name"] = f_val
            elif ("program" in f_lower or "course" in f_lower or "degree" in f_lower) and summary["course"] == "N/A":
                summary["course"] = f_val
            elif ("institute" in f_lower or "college" in f_lower) and summary["institute"] == "N/A":
                summary["institute"] = f_val
            elif ("mobile" in f_lower or "phone" in f_lower or "contact" in f_lower) and summary["contact_mobile"] == "N/A":
                summary["contact_mobile"] = f_val
            elif "domicile" in f_lower and summary["domicile"] == "N/A":
                summary["domicile"] = f_val
            elif "state" in f_lower and summary["domicile"] == "N/A":
                summary["domicile"] = f_val

        log(f"Successfully processed student dossier: {len(all_sections_data)} fields extracted.")
        return {
            "success": True,
            "enrollment_no": enrollment_no,
            "dob": dob,
            "excel_path": excel_path,
            "filename": filename,
            "total_fields": len(all_sections_data),
            "summary": summary,
            "sections": sections_grouped,
            "error": None
        }

    except Exception as exc:
        err_msg = str(exc)
        logger.error(f"Error during scraping: {err_msg}", exc_info=True)
        return {
            "success": False,
            "enrollment_no": enrollment_no,
            "dob": dob,
            "excel_path": None,
            "filename": None,
            "total_fields": 0,
            "summary": {},
            "sections": {},
            "error": err_msg
        }

    finally:
        if driver:
            try:
                driver.quit()
                logger.info("Closed Chrome driver instance.")
            except Exception:
                pass
            driver = None
        cleanup_browser_processes()
        import gc
        gc.collect()
